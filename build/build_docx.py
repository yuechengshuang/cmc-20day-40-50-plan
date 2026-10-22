# -*- coding: utf-8 -*-
"""把 content/*.txt 组装成最终 Word 文档《CMC 20天40-50分基础冲刺计划》。
用法：python build_docx.py
"""
import os
import re
import sys
import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from content_parser import parse_file, FIELDS, FIELD_LABEL
from latex_omml import omml_element, LatexError

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT = os.path.join(ROOT, 'content')
OUT = os.path.join(ROOT, 'output', 'CMC_20天40-50分基础冲刺计划.docx')

MATH_FAILS = []

# ---------------------------------------------------------------- 文档顺序
DOC_ORDER = [
    ('prose', 'my/ch00_usage.txt'),
    ('prose', 'my/ch01_analysis.txt'),
    ('prose', 'my/ch02_plan.txt'),
    ('prose', 'my/ch03_formulas.txt'),
    ('problems', ['ch04a.txt', 'ch04b.txt']),
    ('problems', ['ch05a.txt', 'ch05b.txt']),
    ('problems', ['ch06a.txt', 'ch06b.txt']),
    ('problems', ['ch07a.txt', 'ch07b.txt']),
    ('problems', ['ch08a.txt', 'ch08b.txt']),
    ('problems', ['ch09a.txt', 'ch09b.txt']),
    ('problems', ['ch10a.txt', 'ch10b.txt']),
    ('problems', ['ch11a.txt', 'ch11b.txt']),
    ('problems', ['ch12a.txt', 'ch12b.txt']),
    ('prose', 'my/ch13_mixed.txt'),
    ('mock', ['mock1a.txt', 'mock1b.txt']),
    ('mock', ['mock2a.txt', 'mock2b.txt']),
    ('mock', ['mock3a.txt', 'mock3b.txt']),
    ('prose', 'my/ch14_strategy.txt'),
    ('prose', 'my/ch15_errorbook.txt'),
    ('prose', 'my/appendix.txt'),
]

SOLUTION_FIELDS = ['POINT', 'IDENTIFY', 'FORMULAS', 'IDEA', 'SOLUTION',
                   'ANSWER', 'PITFALL', 'TIPS', 'SUMMARY']

PARA_MARK = re.compile(r'^(\d+[.、]|[-•·]|[①②③④⑤⑥⑦⑧⑨⑩]|（\d+）|\(\d+\)|[a-d][)）])')


# ---------------------------------------------------------------- XML 小工具
def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def left_bar(paragraph, color='9CB8DC', size=18):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), color)
    bdr.append(left)
    pPr.append(bdr)


def box(paragraph, fill=None, color='9CB8DC'):
    if fill:
        shade(paragraph, fill)
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '6')
        e.set(qn('w:space'), '6')
        e.set(qn('w:color'), color)
        bdr.append(e)
    pPr.append(bdr)


def set_run_font(run, ascii_font='Times New Roman', ea='宋体', size=None,
                 bold=None, color=None):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:ascii'), ascii_font)
    rf.set(qn('w:hAnsi'), ascii_font)
    rf.set(qn('w:eastAsia'), ea)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


# ---------------------------------------------------------------- 富文本渲染
TOKEN = re.compile(r'(\*\*.+?\*\*|\$\$.+?\$\$|\$.+?\$)', re.S)


def add_math(p, latex, display=False):
    try:
        p._p.append(omml_element(latex, display=display))
    except (LatexError, Exception) as e:   # 兜底：原样输出，绝不丢内容
        MATH_FAILS.append((latex[:60], str(e)[:60]))
        r = p.add_run(latex)
        set_run_font(r, 'Cambria Math', '宋体')


def fill_rich(p, text, base_bold=False, ea='宋体'):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith('$$') and tok.endswith('$$') and len(tok) > 4:
            add_math(p, tok[2:-2], display=False)
        elif tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            add_math(p, tok[1:-1], display=False)
        elif tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            fill_rich(p, tok[2:-2], base_bold=True, ea=ea)
        else:
            r = p.add_run(tok)
            set_run_font(r, ea=ea, bold=base_bold)


def join_lines(lines):
    out = ''
    for ln in lines:
        if not out:
            out = ln
            continue
        a, b = out[-1], ln[0]
        cjk = lambda ch: '\u3000' <= ch <= '\u9fff' or '\uff00' <= ch <= '\uffef'
        out = out + ('' if (cjk(a) and cjk(b)) else ' ') + ln
    return out


def split_paragraphs(text):
    out = []
    buf = []

    def flush():
        if buf:
            out.append(join_lines(buf))
            del buf[:]

    for ln in text.split('\n'):
        s = ln.strip()
        if not s:
            flush()
            continue
        if (s.startswith('$$') and s.endswith('$$')) or PARA_MARK.match(s):
            flush()
            out.append(s)
            continue
        buf.append(s)
    flush()
    return out


def add_heading_rich(doc, text, level=1):
    """带公式的标题：$...$ 会渲染成 Word 原生公式。"""
    p = doc.add_heading('', level=level)
    fill_rich(p, text, ea='黑体')
    return p


def add_labeled(doc, label, text, fill=None, bar=None, label_color=None,
                indent=0.0, size=None):
    paras = split_paragraphs(text)
    if not paras:
        paras = ['']
    first = True
    for t in paras:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.12
        if indent:
            pf.left_indent = Cm(indent)
        if bar:
            left_bar(p, bar)
        if fill:
            shade(p, fill)
        if t.startswith('$$') and t.endswith('$$') and len(t) > 4:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_math(p, t[2:-2], display=True)
        else:
            if first and label:
                r = p.add_run(label)
                set_run_font(r, ea='黑体', bold=True, size=size)
                if label_color:
                    r.font.color.rgb = label_color
            fill_rich(p, t)
            if size:
                for r in p.runs:
                    if r.font.size is None:
                        r.font.size = Pt(size)
        first = False
    return paras


def add_display_math(doc, latex, space=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(space)
    add_math(p, latex, display=True)


def split_table_row(line):
    """按 | 切分表格行，但公式内部的 | 不切分（如 $\\left|\\frac ab\\right|$）。"""
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    cells, buf, inmath = [], [], False
    for ch in s:
        if ch == '$':
            inmath = not inmath
            buf.append(ch)
        elif ch == '|' and not inmath:
            cells.append(''.join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    cells.append(''.join(buf).strip())
    return cells


def add_table(doc, rows, widths=None, header=True):
    ncols = max(len(r) for r in rows)
    rows = [list(r) + [''] * (ncols - len(r)) for r in rows]
    t = doc.add_table(rows=0, cols=ncols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if ri == 0 and header:
                r = p.add_run(str(val))
                set_run_font(r, ea='黑体', bold=True, size=9.5)
                shade(p, 'DCE6F1')
            else:
                fill_rich(p, str(val))
                for r in p.runs:
                    if r.font.size is None:
                        set_run_font(r, size=9.5)
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    return t


# ---------------------------------------------------------------- 正文块渲染
def render_prose(doc, path):
    if not os.path.exists(path):
        print('  !! 缺少文件：%s' % path)
        return
    with open(path, encoding='utf-8-sig') as f:
        text = f.read().replace('\r\n', '\n').replace('\r', '\n')
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith('### '):
            add_heading_rich(doc, s[4:], 3)
        elif s.startswith('## '):
            add_heading_rich(doc, s[3:], 2)
        elif s.startswith('# '):
            add_heading_rich(doc, s[2:], 1)
        elif s.startswith('$$') and s.endswith('$$') and len(s) > 4:
            add_display_math(doc, s[2:-2])
        elif s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = split_table_row(lines[i])
                if not re.match(r'^[-: ]+$', ''.join(row)):
                    rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(1)
            fill_rich(p, s[2:])
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0.0)
            fill_rich(p, s)
        i += 1


def render_problem(doc, prob, show_stem=True, index=None):
    # 题头
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    shade(p, 'DCE6F1')
    star = prob.star + ('★' * 0)
    title = '题 %s' % prob.id
    if index is not None:
        title = '第 %d 题（%s）' % (index, prob.id)
    head = '%s　【%s】　难度：%s　建议用时：%s　重要程度：%s' % (
        title, prob.type, prob.star, prob.time, prob.imp)
    r = p.add_run(head)
    set_run_font(r, ea='黑体', bold=True, size=10.5)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    shade(p2, 'DCE6F1')
    r = p2.add_run('【知识点】')
    set_run_font(r, ea='黑体', bold=True, size=9.5)
    r = p2.add_run(prob.kp)
    set_run_font(r, size=9.5)

    if show_stem:
        add_labeled(doc, '【题目】', prob.fields.get('STEM', ''), fill='F7F9FC',
                    label_color=RGBColor(0x1F, 0x3B, 0x63))
    for f in SOLUTION_FIELDS:
        label = '【%s】' % FIELD_LABEL[f]
        if f == 'ANSWER':
            add_labeled(doc, label, prob.fields.get(f, ''), fill='FFF6E5',
                        label_color=RGBColor(0xB0, 0x5A, 0x00))
        elif f == 'SUMMARY':
            add_labeled(doc, label, prob.fields.get(f, ''), bar='C0504D',
                        label_color=RGBColor(0x8B, 0x1A, 0x1A))
        elif f in ('PITFALL', 'TIPS'):
            add_labeled(doc, label, prob.fields.get(f, ''), bar='9CB8DC')
        else:
            add_labeled(doc, label, prob.fields.get(f, ''), bar='B7C9E2')


def render_problem_chapter(doc, ch):
    add_heading_rich(doc, ch.title, 1)
    if ch.guide:
        p = doc.add_paragraph()
        r = p.add_run('本章拿分指南')
        set_run_font(r, ea='黑体', bold=True, size=12, color=RGBColor(0x1F, 0x3B, 0x63))
        shade(p, 'EAF1F8')
        add_labeled(doc, '', ch.guide, bar='1F3B63')

    # 本章题型速查表：题号 → 难度 → 母题口诀
    rows = [['题型', '题号范围', '难度', '母题口诀（看到什么 → 想到什么）']]
    for sec in ch.sections:
        if not sec.problems:
            continue
        ids = '%s ~ %s' % (sec.problems[0].id, sec.problems[-1].id)
        dist = {}
        for p in sec.problems:
            dist[p.star] = dist.get(p.star, 0) + 1
        rows.append([sec.full_title, ids,
                     ' '.join('%s×%d' % (k, dist[k]) for k in sorted(dist)),
                     sec.problems[0].fields.get('SUMMARY', '')])
    if len(rows) > 1:
        p = doc.add_paragraph()
        r = p.add_run('本章题型速查表（做完本章后，只看这一张表回忆方法）')
        set_run_font(r, ea='黑体', bold=True, size=11,
                     color=RGBColor(0x8B, 0x1A, 0x1A))
        shade(p, 'FFF2F2')
        add_table(doc, rows, widths=[3.9, 2.6, 2.2, 8.3])

    for sec in ch.sections:
        add_heading_rich(doc, sec.full_title, 2)
        if sec.method:
            p = doc.add_paragraph()
            r = p.add_run('这一类题我到底应该记住什么？')
            set_run_font(r, ea='黑体', bold=True, size=10.5,
                         color=RGBColor(0x8B, 0x1A, 0x1A))
            shade(p, 'FFF2F2')
            add_labeled(doc, '', sec.method, bar='C0504D')
        for prob in sec.problems:
            render_problem(doc, prob)


def render_mock(doc, ch):
    add_heading_rich(doc, ch.title, 1)
    if ch.info:
        p = doc.add_paragraph()
        r = p.add_run('本卷说明（时间分配 · 难度分析 · 分层得分目标）')
        set_run_font(r, ea='黑体', bold=True, size=11,
                     color=RGBColor(0x1F, 0x3B, 0x63))
        shade(p, 'EAF1F8')
        add_labeled(doc, '', ch.info, bar='1F3B63')

    probs = ch.problems
    fill = [p for p in probs if p.type == '填空']
    solve = [p for p in probs if p.type == '解答']
    total = sum(int(re.sub(r'\D', '', p.attrs.get('score', '0')) or 0) for p in probs)

    add_heading_rich(doc, '一、试卷部分（建议限时 150 分钟，满分 %d 分）' % total, 2)
    n = 0
    if fill:
        fs = sum(int(re.sub(r'\D', '', p.attrs.get('score', '0')) or 0) for p in fill)
        p = doc.add_paragraph()
        r = p.add_run('一、填空题（每小题 %d 分，共 %d 分）'
                      % (int(re.sub(r'\D', '', fill[0].attrs.get('score', '6')) or 6), fs))
        set_run_font(r, ea='黑体', bold=True, size=11)
        for prob in fill:
            n += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run('%d. ' % n)
            set_run_font(r, bold=True)
            fill_rich(p, prob.fields.get('STEM', ''))
    if solve:
        ss = sum(int(re.sub(r'\D', '', p.attrs.get('score', '0')) or 0) for p in solve)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run('二、解答题（每小题 %d 分，共 %d 分）'
                      % (int(re.sub(r'\D', '', solve[0].attrs.get('score', '14')) or 14), ss))
        set_run_font(r, ea='黑体', bold=True, size=11)
        for prob in solve:
            n += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run('%d. ' % n)
            set_run_font(r, bold=True)
            fill_rich(p, prob.fields.get('STEM', ''))

    add_heading_rich(doc, '二、参考答案速查', 2)
    rows = [['题号', '题型', '分值', '难度', '答案']]
    for prob in probs:
        rows.append([prob.id, prob.type, prob.attrs.get('score', ''), prob.star,
                     prob.fields.get('ANSWER', '')])
    add_table(doc, rows, widths=[1.6, 1.3, 1.1, 1.5, 8.5])

    add_heading_rich(doc, '三、详细解析', 2)
    for i, prob in enumerate(probs, 1):
        render_problem(doc, prob, show_stem=True, index=i)


# ---------------------------------------------------------------- 文档骨架
def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    for i, (size, ea) in enumerate([(18, '黑体'), (15, '黑体'), (12.5, '黑体'),
                                    (11, '黑体')]):
        st = doc.styles['Heading %d' % (i + 1)]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        rf.set(qn('w:eastAsia'), ea)
        st.paragraph_format.space_before = Pt(12 if i < 2 else 8)
        st.paragraph_format.space_after = Pt(6)

    # 页脚页码
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    set_run_font(r, size=9)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r._element.addnext(fld)


def cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('CMC 20天 40–50分\n基础冲刺计划')
    set_run_font(r, ea='黑体', bold=True, size=30, color=RGBColor(0x1F, 0x3B, 0x63))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('中国大学生数学竞赛（非数学专业 A 类 · 初赛）')
    set_run_font(r, ea='黑体', size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('题库 · 零基础详解 · 母题变式体系 · 3套模拟卷 · 考场策略')
    set_run_font(r, ea='宋体', size=12, color=RGBColor(0x60, 0x60, 0x60))

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('副标题：适用于当前约 20 分、基础薄弱的学生')
    set_run_font(r, ea='楷体', size=13, color=RGBColor(0x8B, 0x1A, 0x1A))

    for _ in range(2):
        doc.add_paragraph()
    rows = [
        ['适用对象', '计算机等理工类专业本科生（CMC 非数学专业 A 类）'],
        ['考试内容', '初赛：高等数学（150 分钟，满分 100 分）'],
        ['卷面结构', '填空题 5 题 × 6 分 + 解答题 5 题 × 14 分'],
        ['当前水平', '真实成绩约 20 分，基础知识大量遗忘'],
        ['训练目标', '20 天稳定冲击 40–50 分（不追求高分与一等奖）'],
        ['训练原则', '基础题 > 常规题 > 中档题 > 难题'],
        ['总题量', '专项训练题 + 模拟卷，全部配零基础详细解析'],
        ['制定日期', datetime.date.today().strftime('%Y 年 %m 月 %d 日')],
    ]
    add_table(doc, rows, widths=[3.2, 12.2], header=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render_toc(doc, entries):
    add_heading_rich(doc, '目　录', 1)
    for level, text in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.0 if level == 1 else 0.8)
        fill_rich(p, text, base_bold=(level == 1), ea='黑体' if level == 1 else '宋体')


# ---------------------------------------------------------------- main
def main():
    doc = Document()
    setup_document(doc)

    loaded = {}

    def get(name):
        if name not in loaded:
            path = os.path.join(CONTENT, name)
            if not os.path.exists(path):
                loaded[name] = None
            else:
                loaded[name] = parse_file(path)
        return loaded[name]

    # 预扫描，生成目录
    toc_entries = []
    plan = []
    for kind, spec in DOC_ORDER:
        if kind in ('prose',):
            path = os.path.join(CONTENT, spec)
            if not os.path.exists(path):
                plan.append(('missing', spec))
                continue
            with open(path, encoding='utf-8-sig') as f:
                first = ''
                for ln in f:
                    if ln.strip().startswith('# '):
                        first = ln.strip()[2:]
                        break
            toc_entries.append((1, first.lstrip('\ufeff') or spec))
            plan.append(('prose', spec))
        else:
            names = spec
            chs = [get(n) for n in names]
            chs = [c for c in chs if c]
            if not chs:
                plan.append(('missing', names[0]))
                continue
            base = chs[0]
            for c in chs[1:]:
                base.sections.extend(c.sections)
                if not base.guide and c.guide:
                    base.guide = c.guide
                if c.info:
                    base.info = (base.info + '\n\n' + c.info).strip()
                if not base.info and c.info:
                    base.info = c.info
            toc_entries.append((1, base.title or names[0]))
            if kind == 'problems':
                for sec in base.sections:
                    toc_entries.append((2, '　'.join([sec.full_title])))
            elif kind == 'mock':
                toc_entries.append((2, '试卷 · 参考答案 · 详细解析'))
            plan.append((kind, base))

    cover(doc)
    render_toc(doc, toc_entries)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    n_prob = 0
    for kind, data in plan:
        if kind == 'missing':
            print('  !! 跳过缺失内容：%s' % data)
            continue
        if kind == 'prose':
            render_prose(doc, os.path.join(CONTENT, data))
        elif kind == 'problems':
            render_problem_chapter(doc, data)
            n_prob += len(data.problems)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == 'mock':
            render_mock(doc, data)
            n_prob += len(data.problems)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print('已生成：%s' % OUT)
    print('题目总数：%d' % n_prob)
    if MATH_FAILS:
        print('公式转换失败 %d 处（已用原文兜底）：' % len(MATH_FAILS))
        for f, e in MATH_FAILS[:20]:
            print('   -', f, '|', e)
    else:
        print('公式转换：全部成功')


if __name__ == '__main__':
    main()