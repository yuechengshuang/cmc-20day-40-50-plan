# -*- coding: utf-8 -*-
"""最小样例 docx：验证公式在 Word/LibreOffice 中的渲染"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from latex_omml import omml_element

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   'output', 'test_formula.docx')

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_heading('公式渲染测试', level=1)
samples = [
    r"\lim_{x\to 0}\frac{\sin x}{x}=1",
    r"\int_0^1 x^2\,dx=\frac13",
    r"f(x)=\begin{cases} x^2, & x\geq 0 \\ -x, & x<0 \end{cases}",
    r"\sum_{n=1}^{\infty}\frac{1}{n^2}=\frac{\pi^2}{6}",
    r"\left.\frac{dy}{dx}\right|_{x=1}=2",
    r"\text{原式}=\lim_{x\to 0}\frac{1-\cos x}{x^2}=\frac12",
    r"\int_{-\infty}^{+\infty}e^{-x^2}\,dx=\sqrt{\pi}",
    r"\lim_{x\to\infty}\left(1+\frac{a}{x}\right)^{x}=e^{a}",
    r"\frac{\partial^2 z}{\partial x\partial y}=2x+y",
    r"\iint_D (x+y)\,dx\,dy",
]
for s in samples:
    # 行内：文字 + 公式
    p = doc.add_paragraph()
    r = p.add_run('行内公式：')
    p._p.append(omml_element(s, display=False))
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2._p.append(omml_element(s, display=True))

doc.save(OUT)
print('saved', OUT)